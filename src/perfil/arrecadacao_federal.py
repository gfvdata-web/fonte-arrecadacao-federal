"""Etapa E — perfil das tabelas da fonte `arrecadacao_federal`.

Esta fonte é a mais rica do projeto em tabelas de apoio, e é por isso que o mapa de
relacionamentos importa aqui:

- **bruto**    — uma aba anual representativa do XLSX da RFB (formato wide: linhas =
  tributos com o nível codificado no recuo, colunas = JAN…DEZ + TOTAL);
- **tidy**     — `dados/processados/arrecadacao_federal.csv`;
- **auxiliar** — inventário das 41 abas dos dois XLSX, a série do IPCA (SGS 433) usada na
  deflação, e a hierarquia de tributos (auto-relacionamento `tributo_pai` → `tributo`).

Gera `docs/dados/perfil_arrecadacao_federal.json`.
"""
from __future__ import annotations

import json

import openpyxl
import pandas as pd

from src import config
from src.perfil import nucleo

SLUG = "arrecadacao_federal"
LINHA_CABECALHO = 6

_MI = "R$ milhões"

DECLARACAO: dict[str, dict] = {
    "bruto": {
        "granularidade": "uma linha por tributo; uma aba por ano; colunas JAN…DEZ + TOTAL",
        "chave_primaria": ["RECEITAS"],
        "colunas": {
            "RECEITAS": ("chave", None,
                         "Rótulo do tributo; o RECUO do texto codifica o nível hierárquico"),
            **{m: ("medida", _MI, f"Arrecadação de {m} no ano da aba")
               for m in ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN",
                         "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]},
            "TOTAL": ("medida", _MI, "Soma dos doze meses, calculada pela própria RFB"),
        },
    },
    "tidy": {
        "granularidade": "uma linha por mês × tributo",
        "chave_primaria": ["ano_mes", "tributo"],
        "colunas": {
            "ano_mes": ("data", None, "Mês de referência (YYYY-MM)"),
            "tributo": ("dimensao", None, "Identificador do tributo, já desambiguado"),
            "rotulo": ("texto", None, "Rótulo de exibição, como publicado pela RFB"),
            "tributo_pai": ("chave", None,
                            "Tributo imediatamente acima na hierarquia; nulo no nível 1"),
            "nivel": ("dimensao", None, "Profundidade na hierarquia (1, 2 ou 3)"),
            "tipo": ("dimensao", None,
                     "componente / detalhe / agregado — o filtro que evita dupla contagem"),
            "valor": ("medida", _MI, "Valor arrecadado a preços correntes"),
            "valor_constante": ("medida", _MI,
                                "Valor deflacionado pelo IPCA a preços do último mês da série"),
        },
    },
    "aux_abas": {
        "granularidade": "uma linha por aba dos dois arquivos XLSX baixados",
        "chave_primaria": ["arquivo", "aba"],
        "colunas": {
            "arquivo": ("dimensao", None, "XLSX de origem"),
            "aba": ("chave", None, "Nome da aba (o ano, quando a aba é anual)"),
            "ano": ("data", None, "Ano da aba; nulo quando a aba cobre um intervalo"),
            "n_linhas": ("medida", "linhas", "Linhas de dado na aba"),
            "n_colunas": ("medida", "colunas", "Colunas na aba"),
            "tratada": ("dimensao", None, "Se a Etapa 3 converte esta aba para o tidy"),
        },
    },
    "aux_ipca": {
        "granularidade": "uma linha por mês",
        "chave_primaria": ["ano_mes"],
        "colunas": {
            "ano_mes": ("data", None, "Mês de referência (YYYY-MM)"),
            "ipca_variacao_pct": ("medida", "% no mês", "IPCA mensal, série 433 do SGS/BCB"),
        },
    },
    "aux_hierarquia": {
        "granularidade": "uma linha por tributo distinto",
        "chave_primaria": ["tributo"],
        "colunas": {
            "tributo": ("chave", None, "Identificador do tributo"),
            "rotulo": ("texto", None, "Rótulo de exibição"),
            "tributo_pai": ("chave", None, "Auto-relacionamento; nulo no nível 1"),
            "nivel": ("dimensao", None, "Profundidade na hierarquia"),
            "tipo": ("dimensao", None, "componente / detalhe / agregado"),
            "meses": ("medida", "meses", "Meses em que o tributo aparece no tidy"),
            "primeiro_mes": ("data", None, "Primeiro mês com valor diferente de zero"),
            "ultimo_mes": ("data", None, "Último mês com valor diferente de zero"),
        },
    },
}


def _pasta() -> "object":
    pasta = config.DIR_BRUTOS / SLUG
    if not (pasta / "_metadados.json").exists():
        raise FileNotFoundError(
            f"Bruto ausente: {pasta}. Rode a Etapa 2 (coleta) antes da Etapa E."
        )
    return pasta


def _aba_representativa(caminho, aba: str) -> pd.DataFrame:
    """Lê uma aba anual como ela é — sem desfazer o formato wide nem o recuo."""
    ws = openpyxl.load_workbook(caminho, read_only=True, data_only=True)[aba]
    linhas = list(ws.iter_rows(min_row=LINHA_CABECALHO, values_only=True))
    cabecalho = [str(c).strip() if c is not None else f"col{i}"
                 for i, c in enumerate(linhas[0])]
    dados = [linha for linha in linhas[1:] if linha[0] is not None]
    return pd.DataFrame(dados, columns=cabecalho)


def _inventario_abas(pasta, cfg: dict) -> pd.DataFrame:
    linhas = []
    for arq in cfg["arquivos"]:
        caminho = pasta / arq["nome"]
        if not caminho.exists():
            continue
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        for aba in wb.sheetnames:
            ws = wb[aba]
            linhas.append({
                "arquivo": arq["nome"],
                "aba": aba,
                "ano": aba if aba.isdigit() else None,
                "n_linhas": int(ws.max_row or 0) - LINHA_CABECALHO,
                "n_colunas": int(ws.max_column or 0),
                "tratada": "sim" if (arq["tratar"] and aba.isdigit()) else "nao",
            })
        wb.close()
    return pd.DataFrame(linhas)


def _ipca(pasta) -> pd.DataFrame:
    registros = json.loads((pasta / "ipca_sgs433.json").read_text(encoding="utf-8"))
    df = pd.DataFrame(registros)
    return pd.DataFrame({
        "ano_mes": df["data"].str.slice(6, 10) + "-" + df["data"].str.slice(3, 5),
        "ipca_variacao_pct": pd.to_numeric(df["valor"], errors="coerce"),
    })


def _hierarquia(tidy: pd.DataFrame) -> pd.DataFrame:
    ativos = tidy[tidy["valor"] != 0]
    agg = (ativos.groupby("tributo")["ano_mes"].agg(["min", "max"])
           .rename(columns={"min": "primeiro_mes", "max": "ultimo_mes"}))
    base = (tidy.groupby("tributo")
            .agg(rotulo=("rotulo", "first"), tributo_pai=("tributo_pai", "first"),
                 nivel=("nivel", "first"), tipo=("tipo", "first"), meses=("ano_mes", "nunique"))
            .join(agg).reset_index())
    return base[["tributo", "rotulo", "tributo_pai", "nivel", "tipo",
                 "meses", "primeiro_mes", "ultimo_mes"]]


def perfilar(slug: str = SLUG) -> dict:
    cfg = config.fonte(slug)
    pasta = _pasta()
    meta_bruto = json.loads((pasta / "_metadados.json").read_text(encoding="utf-8"))

    tidy = pd.read_csv(config.DIR_PROCESSADOS / f"{slug}.csv", dtype={"ano_mes": str})
    abas = _inventario_abas(pasta, cfg)
    ipca = _ipca(pasta)
    hier = _hierarquia(tidy)

    arq_tratado = next(a["nome"] for a in cfg["arquivos"] if a["tratar"])
    ano_ref = str(tidy["ano_mes"].max()[:4])
    bruto = _aba_representativa(pasta / arq_tratado, ano_ref)

    t_bruto = nucleo.perfilar_tabela(
        bruto, id="bruto", camada="bruto",
        nome=f"XLSX da RFB — aba {ano_ref} (representativa das {int((abas['tratada'] == 'sim').sum())} abas tratadas)",
        arquivo=f"dados/brutos/{slug}/{arq_tratado} → aba {ano_ref}",
        formato="XLSX — uma aba por ano, formato wide",
        origem=cfg["url"],
        declaracao=DECLARACAO["bruto"],
        alertas_extra=[
            f"coletado em {meta_bruto.get('coletado_em', '?')}",
            "o NÍVEL hierárquico não é uma coluna: está codificado no recuo do texto da "
            "coluna RECEITAS (0, 2-3 ou 4 espaços). A Etapa 3 é quem o transforma em dado",
            "a coluna TOTAL é soma da própria RFB e não entra no tidy — o tidy guarda os "
            "doze meses e recalcula quando precisa",
            "cada ano é uma aba: o rol de tributos muda entre abas conforme tributos são "
            "criados e extintos, então nem toda aba tem as mesmas linhas",
        ],
    )
    t_tidy = nucleo.perfilar_tabela(
        tidy, id="tidy", camada="tidy",
        nome="Tidy do projeto — 32 abas anuais empilhadas e deflacionadas",
        arquivo=f"dados/processados/{slug}.csv",
        formato="CSV UTF-8",
        origem="Etapa 3 (src/tratamento/arrecadacao_federal.py)",
        declaracao=DECLARACAO["tidy"],
        alertas_extra=[
            "os nulos de `tributo_pai` são ESTRUTURAIS: são exatamente as linhas de nível 1, "
            "que não têm pai. Não é dado faltante",
        ],
    )
    t_abas = nucleo.perfilar_tabela(
        abas, id="aux_abas", camada="auxiliar",
        nome="Inventário das abas dos dois XLSX baixados",
        arquivo=f"dados/brutos/{slug}/*.xlsx",
        formato="tabela derivada",
        origem="Etapa E",
        declaracao=DECLARACAO["aux_abas"],
        alertas_extra=[
            "o arquivo de 1970-1993 é coletado para rastreabilidade mas NÃO é tratado: "
            "mistura quatro moedas (Cr$, Cz$, NCz$, CR$) e traz 1970-1985 em base anual",
        ],
    )
    t_ipca = nucleo.perfilar_tabela(
        ipca, id="aux_ipca", camada="auxiliar",
        nome="IPCA mensal (SGS 433) — deflator da série",
        arquivo=f"dados/brutos/{slug}/ipca_sgs433.json",
        formato="JSON (API SGS/BCB)",
        origem=cfg.get("url_ipca", "https://api.bcb.gov.br/dados/serie/bcdata.sgs.433/dados"),
        declaracao=DECLARACAO["aux_ipca"],
        alertas_extra=[
            "a série do IPCA vai além da série de arrecadação — o deflator existe para meses "
            "que a RFB ainda não publicou",
        ],
    )
    t_hier = nucleo.perfilar_tabela(
        hier, id="aux_hierarquia", camada="auxiliar",
        nome="Hierarquia de tributos (auto-relacionamento pai → filho)",
        arquivo="derivada do tidy",
        formato="tabela derivada",
        origem="Etapa E",
        declaracao=DECLARACAO["aux_hierarquia"],
    )

    tidy_ano = pd.DataFrame({"ano": tidy["ano_mes"].str.slice(0, 4)})
    abas_tratadas = abas[abas["tratada"] == "sim"]

    rel_ipca = nucleo.relacionamento(
        tidy, "ano_mes", ipca, "ano_mes", cardinalidade="N:1",
        uso="deflação: valor_constante = valor corrigido pelo IPCA acumulado até o mês-base")
    rel_hier = nucleo.relacionamento(
        tidy, "tributo", hier, "tributo", cardinalidade="N:1",
        uso="resolve o tributo para nível, tipo e pai")
    rel_pai = nucleo.relacionamento(
        hier, "tributo_pai", hier, "tributo", cardinalidade="N:1 (auto-relacionamento)",
        uso="monta a árvore de tributos; as linhas de nível 1 têm pai nulo e ficam fora da contagem")
    rel_aba = nucleo.relacionamento(
        tidy_ano, "ano", abas_tratadas, "aba", cardinalidade="N:1",
        uso="cada ano do tidy veio de uma aba do XLSX")
    rel_aba["de"] = "ano_mes (ano)"

    perfil = nucleo.montar(
        slug=slug,
        fonte_nome=cfg["nome"],
        tabelas=[t_bruto, t_tidy, t_abas, t_ipca, t_hier],
        relacionamentos=[
            nucleo.liga(rel_ipca, "tidy", "aux_ipca"),
            nucleo.liga(rel_hier, "tidy", "aux_hierarquia"),
            nucleo.liga(rel_pai, "aux_hierarquia", "aux_hierarquia"),
            nucleo.liga(rel_aba, "tidy", "aux_abas"),
        ],
        cobertura=nucleo.cobertura_temporal(tidy),
    )
    nucleo.salvar(perfil, config.DIR_PUBLICADOS / f"perfil_{slug}.json")
    return perfil


if __name__ == "__main__":
    perfilar()
